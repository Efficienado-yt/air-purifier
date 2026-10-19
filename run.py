"""Build every table of the air box comparison into one Excel workbook and a markdown summary.

Usage:
    python3 run.py

Writes airbox_results.xlsx and airbox_report.md next to this file.  Edit params.json (or
fans_data.py) to change any input, then rerun.  A full run enumerates about 106,000 builds
and takes about a minute.
"""
import os
import time

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from fans_data import REF_BOX_AS_BOUGHT, REF_BOX_TODAY
from model import (FANS, HERE, P, all_builds, apply_swap, appliance, cost_per_cfm, fan_cost, filter_set_cost,
                   filter_spec, kit, optimize_swap, room_plan, tau_months)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
SCENARIOS = ["average", "poor", "smoke_month"]
CAPS = P["noise"]["caps_db_3ft"]
OFFSET = P["noise"]["bench_to_3ft_offset_db"]
CHART_DB_MAX = 50
REF_FILTERS = ("16x25x1 MERV13", "16x25x1 Filtrete 1900 (reference box)")
SWAP_CURVE_MONTHS = [1, 2, 3, 6, 9, 12, 18, 24, 36, 48, 60]

MONEY2 = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
DB = "0.0"
F1 = "0.0"

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
REF_FILL = PatternFill("solid", fgColor="FFF2CC")
APP_FILL = PatternFill("solid", fgColor="E2EFDA")

LEGEND = ("Yellow rows = the Efficienado Reference Box (3 NF-P12 + NF-A14; noise measured at the full and half dial, other steps "
          "by fan law; steps are percent of full speed on its voltage dial; airflow from spec), as bought (three fans at $14.95, "
          "spare NF-A14 already on hand) and at the 2026-09-07 fan price (four NF-P12 at $16.45). Green rows = appliances and kits. "
          "Rows marked EST use a spec-sheet fan with no bench data and are never counted as best DIY.")


# ---------------------------------------------------------------------------
# Row helpers
# ---------------------------------------------------------------------------
def is_reference_box(r):
    """The measured Reference Box rows only; the four-NF-P12 recipe is an EST row, not the Reference Box."""
    return r.get("fan") in (REF_BOX_AS_BOUGHT, REF_BOX_TODAY)


def reference_box_rows(rows, cap=None):
    return sorted([r for r in rows if is_reference_box(r) and r["N"] == 4 and r["k"] == 4 and r["filt"] in REF_FILTERS
                   and (cap is None or r["db"] <= cap)], key=lambda r: r["ppc"])


def diy_ok(r):
    """Eligible for 'best DIY': bench-tested fan (not EST) and not the Reference Box."""
    return not r["est"] and not is_reference_box(r)


def label(r):
    if r["kind"] in ("Appliance", "Kit (modeled)"):
        return r["fan"]
    unit = "% speed" if (r.get("est") or is_reference_box(r)) else "%"   # EST and Reference Box steps are percent of full speed
    return f"{r['N']} x {r['fan']} at {r['pwm']}{unit} on {r['k']} x {r['filt']}"


def noise_basis(r):
    if is_reference_box(r):
        return "MEASURED 3 ft" if r["pwm"] in FANS[REF_BOX_AS_BOUGHT]["measured_steps"] else "fan law from measured full dial"
    if r["kind"] == "Appliance":
        return "vendor lab (not HouseFresh)" if "vendor lab" in r["fan"] else "HouseFresh 3 ft"
    return "EST (spec-sheet fan)" if r["est"] else f"bench per fan + 10 log N + {OFFSET:g} dB"


def swap_txt(r):
    return f"{r['life']:g}" if r.get("swaps") else "none"


# (header, value function, number format)
COLS = [
    ("Option", label, None),
    ("Kind", lambda r: r["kind"], None),
    ("Media", lambda r: r["ftype"], None),
    ("CADR new (cfm)", lambda r: round(r["cadr"]), None),
    ("Avg CADR over life (cfm)", lambda r: round(r["avg_cadr"]), None),
    ("dB at 3 ft", lambda r: round(r["db"], 1), DB),
    ("Watts", lambda r: round(r["w"], 1), F1),
    ("Buy $", lambda r: round(r["buy"], 2), MONEY0),
    ("Filter set $", lambda r: round(r["set_cost"], 2), MONEY2),
    ("Best swap (months; blank = no swap inside 5 yr)", lambda r: r["life"] if r.get("swaps") else None, None),
    ("Filter sets in 5 yr (incl. those in the box)", lambda r: r.get("swaps", 0) + r.get("sets_in_box", 1), None),
    ("Filters $/yr (later sets / 5)", lambda r: round(r["filters_usd_yr"], 2), MONEY0),
    ("5-yr total $", lambda r: round(r["five"], 2), MONEY0),
    ("5-yr $ per avg cfm", lambda r: round(r["ppc"], 2), MONEY2),
    ("Upfront $ per new cfm", lambda r: round(r["upfront"], 2), MONEY2),
    ("New cfm per W", lambda r: round(r["cfm_per_w"], 1), F1),
    ("Media index (1 in pleat of same face = 1.0)", lambda r: round(r["media"], 1) if r["media"] else "", None),
    ("Face velocity fpm", lambda r: round(r["v"]) if r["v"] else "", None),
    ("Filter push-back Pa", lambda r: round(r["dp"], 1) if r["dp"] is not None else "", F1),
    ("tau (months)", lambda r: round(r["tau"], 1) if r["tau"] else "", F1),
    ("Noise basis", noise_basis, None),
    ("Note", lambda r: r.get("note", ""), None),
]
COL_NAMES = [h for h, _, _ in COLS]


# ---------------------------------------------------------------------------
# Worksheet helpers
# ---------------------------------------------------------------------------
def header(ws, row, headers):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")


def row_fill(r):
    if r.get("kind") in ("Appliance", "Kit (modeled)"):
        return APP_FILL
    if is_reference_box(r):
        return REF_FILL
    return None


def write_table(ws, rows, start_row=1, extra_cols=None, title=None):
    """Write rows with the standard columns (plus any extra leading columns); return (header row, last row)."""
    cols = [(h, f, None) for h, f in (extra_cols or [])] + COLS
    r0 = start_row
    if title:
        ws.cell(row=r0, column=1, value=title).font = Font(bold=True, size=12)
        r0 += 1
    header(ws, r0, [h for h, _, _ in cols])
    for i, r in enumerate(rows, r0 + 1):
        fill = row_fill(r)
        for j, (h, fn, fmt) in enumerate(cols, 1):
            c = ws.cell(row=i, column=j, value=fn(r))
            if fmt:
                c.number_format = fmt
            if fill:
                c.fill = fill
    return r0, r0 + len(rows)


def autosize(ws, widths=None, first=58):
    for j in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 14)
    ws.column_dimensions["A"].width = first


def color_scale(ws, col_letter, first, last):
    ws.conditional_formatting.add(f"{col_letter}{first}:{col_letter}{last}",
                                  ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50,
                                                 mid_color="FFEB84", end_type="max", end_color="F8696B"))


# ---------------------------------------------------------------------------
# Compute
# ---------------------------------------------------------------------------
def compute():
    """Every DIY build, appliance and kit row, per scenario."""
    builds = {sc: all_builds(sc) for sc in SCENARIOS}
    apps = {sc: [appliance(n, sc) for n in P["appliances"] if not n.startswith("_")] for sc in SCENARIOS}
    kits = {sc: [kit(n, sc) for n in P["kits"] if not n.startswith("_")] for sc in SCENARIOS}
    return builds, apps, kits


def best_diy(builds, sc, cap):
    return sorted([r for r in builds[sc] if r["db"] <= cap and diy_ok(r)], key=lambda r: r["ppc"])


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------
def sheet_readme(wb):
    ws = wb.active
    ws.title = "README"
    SC_UG = P["filter_life"]["scenarios_ug_per_m3_indoor_with_cleaner_running"]
    DER = P["appliance_loading_derate_avg_cadr"]
    lines = [
        "Air box comparison, generated by run.py from params.json and fans_data.py.",
        "Every input is in params.json (prices, rate, catch rates, decay, noise offset, rooms). Edit and rerun.",
        "Prices (read 2026-09-07 in a browser at Amazon, Micro Center, Lowe's, Walmart, Target, Newegg, manufacturer stores): fans and filters are priced by the packs a builder would buy (fan_packs, filter_packs in params.json); the cost of N fans or k filters is the cheapest pack combination that yields at least N or k, spares included. Every later filter set is priced the same way.",
        "",
        "Sheets (11):",
        "Summary: for each scenario and noise cap (dB at 3 ft), the six cheapest DIY builds, the Reference Box, and every appliance and kit under the cap, together, sorted by dollars per average delivered cfm. Estimated-fan (EST) rows are excluded from the DIY six.",
        "Chart data: the points behind the Summary scatter (average scenario, 150 cheapest DIY builds plus appliances and kits, all at or under 50 dB).",
        "Rooms: a shortlist (best two builds per media type per cap, the Reference Box, appliances, kits) stacked to 5 air changes an hour in each room, then, per noise cap, the cheapest option over every build.",
        f"Builds average / poor / smoke_month: every DIY build the model enumerated (fan, count, filter, count, speed) with new CADR at or above {P['enumeration']['min_cadr_cfm']} cfm, sorted by dollars per cfm.",
        "Swap curves: dollars per cfm vs swap interval for the shortlist candidates, at three media-sensitivity settings.",
        "Fans, Filters, Params: the inputs.",
        "",
        "Definitions:",
        "Filter push-back: the Nukit MERV 13 polynomial on the face velocity divided by an effective-area factor that grows with pleat depth (1 in 1.0, 2 in and 4 in 1.76; fitted to FilterBoxBuilder's own 20x20x1 vs 20x20x2 readings; see params dp_area_factor_by_depth).",
        f"CADR new = airflow through the filters x build factor ({P['build_factor']} DIY, {P.get('build_factor_kit', P['build_factor'])} sealed kits) x single-pass catch of the media when new. Smoke-class (0.3 to 1 um), modeled, not measured; validated against seven measured PC-fan boxes (params.json, filter_types._validation).",
        "Avg CADR over life = CADR new averaged over the filter's decay across the five years (full swap cycles plus the partial last one), at the swap interval that minimises dollars per cfm (DIY, kits); appliances: CADR new x the HEPA loading derate, or the same electret decay on a yearly swap for MERV 13 appliances (box-fan CR box, Tempest Pro stock).",
        f"dB at 3 ft = Cybenetics chamber dBA for one fan + 10 log(fans) + {OFFSET:g} dB (one anchor: Nukit Tempest). Appliances: HouseFresh measured at 3 ft. Reference Box: measured at 3 ft at the full and half dial (Noise basis says MEASURED on those two steps), fan law on the other steps. Rows marked EST use a fan with no bench data (the NF-P12 redux, from Noctua's spec sheet).",
        "5-yr $ per avg cfm = (purchase, which includes the first filter set + later filter sets at the best swap interval + five years of electricity 24/7) / average CADR.",
        f"Appliances: price includes the first filter set, four more yearly sets billed (three where two sets ship in the box; five where none does, the Tempest Pro stock row; nine 6-month filters for the Blue Pure; IQAir's three parts at their stated lives, see params appliance_options); average CADR = new x {DER['average']} / {DER['poor']} / {DER['smoke_month']} (average / poor / smoke month) for HEPA loading at quiet speed (Huang 2024, Winix C535, rescaled from Turbo; see params.json). Watts held constant for appliances, kits and DIY alike.",
        "DIY and kits: a blank 'Best swap' means the optimizer found no swap worth making inside five years, so 'Filter sets in 5 yr' is 1 and filters $/yr is 0. 'Filters $/yr' = later sets / 5 on every row, so it is comparable across kinds.",
        "Fan steps at or below the Cybenetics chamber floor (6.0 dBA) are excluded; steps just above it (6.5 to 8 dBA) are kept and are close to the chamber's limit. Nothing below about 24 dB at 3 ft is measurable in a real room.",
        "",
        f"Scenarios (indoor particle load with the cleaner running): average {SC_UG['average']} ug/m3, poor {SC_UG['poor']}, smoke month {SC_UG['smoke_month']}.",
        f"Media sensitivity s = {P['filter_life']['decay']['sensitivity_default']} (0 = filter life does not depend on the box's media area, 1 = proportional to media per cfm). Swap curves sheet shows 0, 0.5 and 1.",
        "Media index = face sq ft x a pleat-depth multiplier (1 in = 1.0, 2 in = 2.14, 4 in = 2.97). It is a relative index, not square feet of media (a 1 in pleat holds about 2 sq ft of media per sq ft of face).",
        LEGEND,
        "ASSUMPTIONS flagged in params.json: MERV 11 catch 0.85 (solved from CleanAirKits' paired tests, per-box range 0.80 to 0.88; at 0.80 MERV 13 wins every cap instead, so 'MERV 11 vs 13' is a tie inside the model's error, not a ranking), MERV 16 catch and resistance, H11 panel resistance, the bench-to-room offset, kit build factor, the decay floor and tau values (the 'no swap in five years' result flips below a relative floor of 0.376; the 2 inch vs 1 inch choice is inside 1%), the appliance loading derate, the Tempest Pro stock-fan watts.",
    ]
    for i, line in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 160


def sheet_summary(wb, builds, apps, kits):
    ws = wb.create_sheet("Summary")
    ws.cell(row=1, column=1, value=LEGEND).font = Font(italic=True)
    row = 3
    for sc in SCENARIOS:
        for cap in CAPS:
            diy = best_diy(builds, sc, cap)[:6]
            ref = reference_box_rows(builds[sc], cap)[:6]
            ap = sorted([a for a in apps[sc] + kits[sc] if a["db"] <= cap], key=lambda r: r["ppc"])
            rows = sorted(diy + ref + ap, key=lambda r: r["ppc"])
            for r in rows:
                r["_cap"] = cap
                r["_sc"] = sc
            r0, r1 = write_table(ws, rows, row, extra_cols=[("Scenario", lambda r: r["_sc"]), ("Noise cap dB", lambda r: r["_cap"])],
                                 title=f"Scenario {sc}, whole unit at or under {cap} dB at 3 ft")
            row = r1 + 2
    autosize(ws)

    # scatter: $/cfm vs dB for the average scenario, cheapest DIY builds plus appliances and kits, same noise window
    ws_c = wb.create_sheet("Chart data")
    diy_pts = sorted([r for r in builds["average"] if r["db"] <= CHART_DB_MAX and diy_ok(r)], key=lambda r: r["ppc"])[:150]
    other_pts = [r for r in apps["average"] + kits["average"] if r["db"] <= CHART_DB_MAX]
    hdr = ["Option", "dB at 3 ft", "5-yr $ per avg cfm", "Kind"]
    ws_c.append(hdr)
    header(ws_c, 1, hdr)
    for r in diy_pts + other_pts:
        ws_c.append([label(r), round(r["db"], 1), round(r["ppc"], 2), r["kind"]])
    autosize(ws_c)
    ch = ScatterChart()
    ch.title = f"Average air: five-year $ per avg cfm vs noise (150 cheapest DIY builds, appliances and kits, all at or under {CHART_DB_MAX} dB)"
    ch.x_axis.title = "dB at 3 ft"
    ch.y_axis.title = "$ per avg cfm"
    ch.style = 13
    n_diy = len(diy_pts)
    xs = Reference(ws_c, min_col=2, min_row=2, max_row=1 + n_diy)
    ys = Reference(ws_c, min_col=3, min_row=2, max_row=1 + n_diy)
    s1 = Series(ys, xs, title="DIY builds")
    s1.marker.symbol = "circle"
    s1.graphicalProperties.line.noFill = True
    ch.series.append(s1)
    xa = Reference(ws_c, min_col=2, min_row=2 + n_diy, max_row=1 + n_diy + len(other_pts))
    ya = Reference(ws_c, min_col=3, min_row=2 + n_diy, max_row=1 + n_diy + len(other_pts))
    s2 = Series(ya, xa, title="Appliances and kits")
    s2.marker.symbol = "diamond"
    s2.marker.size = 9
    s2.graphicalProperties.line.noFill = True
    ch.series.append(s2)
    ch.width = 28
    ch.height = 14
    ws.add_chart(ch, "Y3")


def shortlist(builds):
    """Best two DIY builds per media type per cap, plus the measured Reference Box steps; deduplicated by label."""
    cur = []
    for ft in P["filter_types"]:
        if ft.startswith("_"):
            continue
        for cap in CAPS:
            cur += sorted([r for r in builds["average"] if r["ftype"] == ft and r["db"] <= cap and diy_ok(r)], key=lambda r: r["ppc"])[:2]
    ref_rows = [r for r in reference_box_rows(builds["average"]) if r["pwm"] in FANS[REF_BOX_AS_BOUGHT]["measured_steps"]]
    return list({label(r): r for r in cur + ref_rows}.values())


def sheet_rooms(wb, builds, apps, kits, cur):
    ws = wb.create_sheet("Rooms")
    ws.cell(row=1, column=1, value=LEGEND).font = Font(italic=True)
    row = 3
    cands = list(cur) + apps["average"] + kits["average"]
    # the Reference Box is a reference (its as-bought fan price includes a spare on hand), shown in the shortlist, not a recipe to win with
    allrows = [x for x in builds["average"] if diy_ok(x)] + apps["average"] + kits["average"]
    for room in P["rooms"]:
        if room.startswith("_"):
            continue
        R = P["rooms"][room]
        ws.cell(row=row, column=1, value=f"{room}: {R['volume_cuft']} cu ft, target {R['target_each']} ACH = {R['target_each'] * R['volume_cuft'] / 60:.0f} cfm of clean air (average-air scenario, average CADR over filter life)").font = Font(bold=True, size=12)
        row += 1
        hdr = ["Option (shortlist)", "Kind", "Units needed", "Delivered cfm (avg)", "ACH reached", "Total dB at 3 ft", "5-yr total $ (all units)",
               "5-yr $ per unit", "Per-unit avg CADR", "Per-unit dB", "Noise basis"]
        header(ws, row, hdr)
        plans = [(r, room_plan(r, room)) for r in cands]
        plans = [(r, pl) for r, pl in plans if pl]
        plans.sort(key=lambda x: x[1]["five"])
        first = row + 1
        for r, pl in plans:
            row += 1
            vals = [label(r), r["kind"], pl["units"], round(pl["cfm"], 1), round(pl["each"], 1), round(pl["db"], 1), round(pl["five"], 2),
                    round(r["five"], 2), round(r["avg_cadr"], 1), round(r["db"], 1), noise_basis(r)]
            fmts = [None, None, None, F1, F1, DB, MONEY0, MONEY0, F1, DB, None]
            fill = row_fill(r)
            for j, (v, fm) in enumerate(zip(vals, fmts), 1):
                c = ws.cell(row=row, column=j, value=v)
                if fm:
                    c.number_format = fm
                if fill:
                    c.fill = fill
        color_scale(ws, "G", first, row)
        color_scale(ws, "F", first, row)
        row += 2
        ws.cell(row=row, column=1, value=f"{room}: cheapest option (searched over every non-EST build, appliances and kits; the Reference Box excluded, see shortlist) whose stacked total stays at or under each noise cap").font = Font(bold=True)
        row += 1
        header(ws, row, ["Noise cap dB at 3 ft", "Option", "Units", "Total dB", "5-yr total $", "Noise basis"])
        allplans = [(rr, room_plan(rr, room)) for rr in allrows]
        allplans = [(r, pl) for r, pl in allplans if pl]
        allplans.sort(key=lambda x: x[1]["five"])
        for cap in CAPS:
            ok = [(r, pl) for r, pl in allplans if pl["db"] <= cap]
            row += 1
            if ok:
                r, pl = ok[0]
                vals = [cap, label(r), pl["units"], round(pl["db"], 1), round(pl["five"], 2), noise_basis(r)]
            else:
                vals = [cap, "nothing reaches the target under this cap", "", "", "", ""]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=row, column=j, value=v)
                if j == 4 and v != "":
                    c.number_format = DB
                if j == 5 and v != "":
                    c.number_format = MONEY0
        row += 3
    autosize(ws, {2: 12, 3: 12, 4: 16, 5: 12, 6: 14, 7: 18, 8: 14, 9: 14, 10: 12, 11: 18})


def sheets_builds(wb, builds):
    for sc in SCENARIOS:
        ws = wb.create_sheet(f"Builds {sc}")
        rows = sorted(builds[sc], key=lambda r: r["ppc"])
        r0, r1 = write_table(ws, rows, 1)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{r1}"
        ppc_col = get_column_letter(COL_NAMES.index("5-yr $ per avg cfm") + 1)
        db_col = get_column_letter(COL_NAMES.index("dB at 3 ft") + 1)
        color_scale(ws, ppc_col, 2, r1)
        color_scale(ws, db_col, 2, r1)
        autosize(ws)


def sheet_swap_curves(wb, cur):
    ws = wb.create_sheet("Swap curves")
    hdr = ["Candidate (average air)", "s", "tau (months)", "floor"] + [f"{L} mo" for L in SWAP_CURVE_MONTHS] + ["best swap (months)", "best $/cfm"]
    ws.append(hdr)
    header(ws, 1, hdr)
    for r in cur:
        ft = P["filter_types"][r["ftype"]]
        for s in [0, 0.5, 1]:
            tau = tau_months(r["ftype"], "average", r["media"], r["q"], s)
            b = optimize_swap(r["buy"], r["set_cost"], r["w"], r["cadr"], tau, ft["decay"]["floor"], ft.get("max_swap_months"))
            ws.append([label(r), s, round(tau, 1), ft["decay"]["floor"]]
                      + [round(cost_per_cfm(L, r["buy"], r["set_cost"], r["w"], r["cadr"], tau, ft["decay"]["floor"])[0], 2) for L in SWAP_CURVE_MONTHS]
                      + [b["L"], round(b["ppc"], 2)])
    autosize(ws, {j: 9 for j in range(2, 20)})


def sheets_inputs(wb):
    ws = wb.create_sheet("Fans")
    hdr = ["Fan", "Size mm", "Max cfm", "Max mm H2O", "Max W", "Max dBA bench (per fan)", "cfm at 25 dBA bench", "Price per fan $ (single)",
           "cfm per $ at 25 dBA", "Price basis", "Data basis", "Packs (count @ price), read 2026-09-07", "Cost for 2 / 4 / 6 / 8 fans"]
    ws.append(hdr)
    header(ws, 1, hdr)
    for n, f in FANS.items():
        c25 = None
        for i in range(len(f["steps"]) - 1):
            if f["db"][i] <= 25 <= f["db"][i + 1]:
                fr = (25 - f["db"][i]) / (f["db"][i + 1] - f["db"][i] or 1)
                c25 = f["cfm"][i] + fr * (f["cfm"][i + 1] - f["cfm"][i])
        basis = "Cybenetics measured"
        if f.get("estimated"):
            basis = "EST: curve and noise from the spec sheet"
        if f.get("measured"):
            basis = "noise MEASURED at 3 ft (per-fan value backed out: minus offset and 10 log 4); curve from spec by fan law; price is 3 x $14.95 / 4 (the NF-A14 was a spare on hand)"
        pk = P.get("fan_packs", {}).get(n, [])
        ws.append([n, f["size"], round(f["cfm"][-1], 1), round(f["mm"][-1], 2), round(f["w"][-1], 2), round(f["db"][-1], 1),
                   round(c25, 1) if c25 else "", round(f["price"], 2), round(c25 / f["price"], 1) if c25 else "", f["pnote"], basis,
                   "; ".join(f"{c} @ ${p:.2f} ({s})" for c, p, s in pk), " / ".join(f"${fan_cost(n, k):.2f}" for k in (2, 4, 6, 8))])
    autosize(ws)

    ws = wb.create_sheet("Filters")
    hdr = ["Filter", "W in", "H in", "Depth in", "$ per filter (legacy single)", "Media type", "Face sq ft", "Media index (1 in = 1.0)", "Source",
           "Packs (count @ price), read 2026-09-07", "Set cost k=1 / 2 / 4"]
    ws.append(hdr)
    header(ws, 1, hdr)
    for n in P["filters"]:
        if n.startswith("_"):
            continue
        w, h, d, price, ft, src = filter_spec(n)
        pk = P.get("filter_packs", {}).get(n, [])
        ws.append([n, w, h, d, price, ft, round(w * h / 144, 2), round(w * h / 144 * P["media_sqft_per_face_sqft_by_depth"][str(int(d))], 2), src,
                   "; ".join(f"{c} @ ${p:.2f} ({s})" for c, p, s in pk), " / ".join(f"${filter_set_cost(n, k):.2f}" for k in (1, 2, 4))])
    ws.append([])
    hdr2 = ["Media type", "Catch new (single pass, smoke class)", "Decay floor (fraction of new)", "tau ref, months, average (reference box)",
            "tau ref, poor", "tau ref, smoke month", "Push-back model", "Max swap (months; blank = none)", "Notes"]
    ws.append(hdr2)
    header(ws, ws.max_row, hdr2)
    for ft, v in P["filter_types"].items():
        if ft.startswith("_"):
            continue
        ws.append([ft, v["catch"], v["decay"]["floor"], v["decay"]["tau_ref_months"]["average"], v["decay"]["tau_ref_months"]["poor"],
                   v["decay"]["tau_ref_months"]["smoke_month"], v["dp"], v.get("max_swap_months"), v["notes"]])
    autosize(ws)

    ws = wb.create_sheet("Params")
    ws.append(["key", "value"])
    header(ws, 1, ["key", "value"])

    def walk(d, pre=""):
        for k, v in d.items():
            if isinstance(v, dict):
                walk(v, pre + k + ".")
            else:
                ws.append([pre + k, str(v)])

    walk(P)
    autosize(ws, {2: 120})


# ---------------------------------------------------------------------------
# Markdown summary
# ---------------------------------------------------------------------------
def markdown_report(builds, apps, kits):
    DER = P["appliance_loading_derate_avg_cadr"]
    half_step = FANS[REF_BOX_AS_BOUGHT]["measured_steps"][0]
    md = [f"# Air box summary (generated {time.strftime('%Y-%m-%d')})", "",
          f"Five-year dollars per average delivered cfm, 24/7, ${P['electricity_rate_usd_per_kwh']}/kWh. DIY noise = bench + {OFFSET:g} dB; "
          f"appliances HouseFresh at 3 ft; the Reference Box measured at 3 ft at the full and half dial (100% and {half_step}% of full speed, "
          f"room floor subtracted), fan law between; its airflow is from spec, not measured. 'Reference Box' = the Efficienado Reference Box, "
          f"3 NF-P12 + NF-A14 on four 16x25x1, best of Nordic Pure MERV 13 and Filtrete 1900 (the filter is named). Appliance avg CADR = new x "
          f"loading derate ({DER['average']} / {DER['poor']} / {DER['smoke_month']}), or the model's electret decay for MERV 13 appliances.", ""]
    for sc in SCENARIOS:
        md.append(f"## {sc}")
        md.append("| Cap dB at 3 ft | Best DIY build | CADR new | avg CADR | dB | W | Buy | swap mo | filters $/yr | 5-yr $/cfm | "
                  "Best appliance under cap | its $/cfm | Best kit under cap | its $/cfm | Reference Box under cap | its $/cfm |")
        md.append("|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|")
        for cap in CAPS:
            d = best_diy(builds, sc, cap)
            a = sorted([x for x in apps[sc] if x["db"] <= cap], key=lambda r: r["ppc"])
            kk = sorted([x for x in kits[sc] if x["db"] <= cap], key=lambda r: r["ppc"])
            h = [x for x in reference_box_rows(builds[sc], cap) if x["fan"] == REF_BOX_AS_BOUGHT]
            ht = [x for x in reference_box_rows(builds[sc], cap) if x["fan"] == REF_BOX_TODAY]
            b = d[0] if d else None
            left = (f"| {cap} | {label(b)} | {b['cadr']:.0f} | {b['avg_cadr']:.0f} | {b['db']:.1f} | {b['w']:.1f} | ${b['buy']:.0f} | {swap_txt(b)} | "
                    f"${b['filters_usd_yr']:.0f} | ${b['ppc']:.2f} |") if b else f"| {cap} | none | | | | | | | | |"
            right = (f" {a[0]['fan'] if a else '-'} | {('$%.2f' % a[0]['ppc']) if a else '-'} | {kk[0]['fan'] if kk else '-'} | "
                     f"{('$%.2f' % kk[0]['ppc']) if kk else '-'} | "
                     f"{(str(h[0]['pwm']) + '% speed, ' + h[0]['filt'] + ', ' + ('%.1f' % h[0]['db']) + ' dB') if h else 'none'} | "
                     f"{('$%.2f as bought, $%.2f at current fan price' % (h[0]['ppc'], ht[0]['ppc'])) if h and ht else '-'} |")
            md.append(left + right)
        md.append("")
    # swap-cap sensitivity: the headline builds re-optimised with the swap forced inside N months
    sens = P["filter_life"]["decay"].get("swap_cap_sensitivity_months", [])
    if sens:
        md.append("## Swap-cap sensitivity (headline DIY build per cap, swap forced to at most N months)")
        md.append("The optimizer often answers no swap inside five years (the boundary of the search) or one swap at 30 months. "
                  "These rows show the same builds with a swap forced inside 24 and 36 months.")
        md.append("")
        md.append("| Scenario | Cap | Build | $/cfm, optimizer | " + " | ".join(f"$/cfm, swap <= {m} mo (swap chosen)" for m in sens) + " |")
        md.append("|---|---|---|---|" + "---|" * len(sens))
        for sc in SCENARIOS:
            for cap in CAPS:
                d = best_diy(builds, sc, cap)
                if not d:
                    continue
                b = d[0]
                ft = P["filter_types"][b["ftype"]]
                cells = []
                for m in sens:
                    rr = dict(b)
                    apply_swap(rr, ft, max_L=m)
                    cells.append(f"${rr['ppc']:.2f} ({rr['life']:g} mo)")
                md.append(f"| {sc} | {cap} | {label(b)} | ${b['ppc']:.2f} ({swap_txt(b)}) | " + " | ".join(cells) + " |")
        md.append("")
    return "\n".join(md)


# ---------------------------------------------------------------------------
def main():
    t0 = time.time()
    print("computing builds...", flush=True)
    builds, apps, kits = compute()
    print(f"{sum(len(v) for v in builds.values())} builds in {time.time() - t0:.0f} s", flush=True)

    wb = Workbook()
    sheet_readme(wb)
    sheet_summary(wb, builds, apps, kits)
    cur = shortlist(builds)
    sheet_rooms(wb, builds, apps, kits, cur)
    sheets_builds(wb, builds)
    sheet_swap_curves(wb, cur)
    sheets_inputs(wb)
    out = os.path.join(HERE, "airbox_results.xlsx")
    wb.save(out)

    md = markdown_report(builds, apps, kits)
    with open(os.path.join(HERE, "airbox_report.md"), "w") as fh:
        fh.write(md)
    print(md)
    print(f"wrote {out} in {time.time() - t0:.0f} s")


if __name__ == "__main__":
    main()
